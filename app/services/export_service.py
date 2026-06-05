from __future__ import annotations

from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.schemas.quotation import QuotationResponse
from app.services.quotation_service import IMPORTANT_NOTES


HEADER_FILL = PatternFill("solid", fgColor="234E52")
SECTION_FILL = PatternFill("solid", fgColor="DDEDEA")
WHITE_FONT = Font(color="F8FBFA", bold=True)
BOLD_FONT = Font(bold=True)
THIN_SIDE = Side(style="thin", color="B7CCC8")
TABLE_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)


def build_quotation_workbook(quotation: QuotationResponse) -> BytesIO:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "报价表"
    worksheet.sheet_view.showGridLines = False

    worksheet.merge_cells("A1:F1")
    worksheet["A1"] = "新案基础报价表"
    worksheet["A1"].font = Font(size=18, bold=True, color="17383B")
    worksheet["A1"].alignment = Alignment(horizontal="center")

    row = 3
    row = _write_section_title(worksheet, row, "基本信息")
    basic_info = [
        ("报价编号", quotation.quotation_no, "报价日期", quotation.created_at.date().isoformat()),
        ("客户名称", quotation.client_name, "客户联系人", quotation.client_contact or "-"),
        ("国家/地区", quotation.country_code, "申请类型", quotation.application_type),
        ("申请途径", quotation.filing_route, "币种", quotation.currency),
        ("顾问", quotation.consultant_name, "报价状态", quotation.status),
    ]
    for left_label, left_value, right_label, right_value in basic_info:
        worksheet.cell(row=row, column=1, value=left_label)
        worksheet.cell(row=row, column=2, value=left_value)
        worksheet.cell(row=row, column=4, value=right_label)
        worksheet.cell(row=row, column=5, value=right_value)
        for column in (1, 2, 4, 5):
            cell = worksheet.cell(row=row, column=column)
            cell.border = TABLE_BORDER
            cell.alignment = Alignment(vertical="center")
        for column in (1, 4):
            worksheet.cell(row=row, column=column).font = BOLD_FONT
            worksheet.cell(row=row, column=column).fill = SECTION_FILL
        row += 1

    row += 1
    row = _write_section_title(worksheet, row, "费用明细")
    headers = ["阶段", "费用项目", "费用类型", "金额", "币种", "费用性质", "备注"]
    _write_table_header(worksheet, row, headers)
    row += 1
    for item in sorted(quotation.items, key=lambda value: value.sort_order):
        values = [
            item.stage,
            item.item_name,
            item.fee_type,
            item.amount,
            item.currency,
            item.cost_nature,
            item.remark,
        ]
        for index, value in enumerate(values, start=1):
            cell = worksheet.cell(row=row, column=index, value=value)
            cell.border = TABLE_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if index == 4:
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal="right")
        row += 1

    row += 1
    row = _write_section_title(worksheet, row, "费用汇总")
    summary_headers = ["费用类别", "金额", "币种", "说明"]
    _write_table_header(worksheet, row, summary_headers)
    row += 1
    summary_rows = [
        ("申请阶段费用", quotation.current_stage_total, quotation.currency, "当前阶段费用"),
        ("后续预估费用", quotation.future_stage_total, quotation.currency, "审查、授权、年费等后续费用"),
        ("全流程预估合计", quotation.total_amount, quotation.currency, "仅供参考"),
    ]
    for label, amount, currency, remark in summary_rows:
        for index, value in enumerate([label, amount, currency, remark], start=1):
            cell = worksheet.cell(row=row, column=index, value=value)
            cell.border = TABLE_BORDER
            if index == 2:
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal="right")
        worksheet.cell(row=row, column=1).font = BOLD_FONT
        if label == "全流程预估合计":
            for column in range(1, 5):
                worksheet.cell(row=row, column=column).font = BOLD_FONT
                worksheet.cell(row=row, column=column).fill = SECTION_FILL
        row += 1

    row += 1
    row = _write_section_title(worksheet, row, "重要说明")
    for note_index, note in enumerate(IMPORTANT_NOTES, start=1):
        worksheet.cell(row=row, column=1, value=f"{note_index}. {note}")
        worksheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        worksheet.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        row += 1

    row += 1
    worksheet.cell(row=row, column=1, value="报价有效期")
    worksheet.cell(row=row, column=2, value="30 天")
    worksheet.cell(row=row, column=1).font = BOLD_FONT

    _apply_layout(worksheet)

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def _write_section_title(worksheet, row: int, title: str) -> int:
    worksheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
    cell = worksheet.cell(row=row, column=1, value=title)
    cell.fill = SECTION_FILL
    cell.font = Font(bold=True, color="17383B")
    cell.alignment = Alignment(vertical="center")
    return row + 1


def _write_table_header(worksheet, row: int, headers: list[str]) -> None:
    for index, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=row, column=index, value=header)
        cell.fill = HEADER_FILL
        cell.font = WHITE_FONT
        cell.border = TABLE_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _apply_layout(worksheet) -> None:
    widths = [16, 28, 14, 14, 12, 14, 34]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    for row in worksheet.iter_rows():
        for cell in row:
            cell.alignment = cell.alignment.copy(vertical="center")
            if isinstance(cell.value, Decimal):
                cell.value = float(cell.value)
    worksheet.freeze_panes = "A9"
