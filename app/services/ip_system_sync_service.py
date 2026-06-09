"""Legacy advanced IP-system import/sync service.

The lightweight query page does not use candidate/review/batch workflows.
"""

from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import re
from datetime import date, datetime
from uuid import uuid4

from openpyxl import load_workbook

from app.db import mysql
from app.schemas.ip_system import (
    IpSystemChangeReview,
    IpSystemImportParseError,
    IpSystemImportRequest,
    IpSystemImportResponse,
    IpSystemReprocessBatchResponse,
    IpSystemSyncBatch,
    IpSystemSyncBatchDetail,
)


IMPORT_FIELDS = [
    "system_code",
    "business_domain",
    "relation_type_code",
    "jurisdiction_name",
    "official_name",
    "official_code",
    "effective_date",
    "expiry_date",
    "is_active",
    "quote_hint_enabled",
    "path_rule_dependency",
    "quote_hint_text",
    "special_statement",
    "source_reference",
    "data_quality_flags",
    "admin_remark",
]


def create_source_snapshot(
    *,
    system_id: str,
    source_config_id: str | None,
    snapshot_type: str,
    raw_snapshot_path: str,
    raw_content: str,
    raw_metadata: dict[str, object] | None,
    captured_by: str,
    remark: str = "",
) -> dict[str, object]:
    snapshot = {
        "snapshot_id": f"ipsnap-{uuid4().hex[:16]}",
        "system_id": system_id,
        "source_config_id": source_config_id,
        "snapshot_type": snapshot_type,
        "raw_snapshot_path": raw_snapshot_path,
        "raw_content_hash": hashlib.sha256(raw_content.encode("utf-8")).hexdigest(),
        "raw_metadata_json": json.dumps(raw_metadata or {}, ensure_ascii=False, default=str),
        "captured_at": datetime.now(),
        "captured_by": captured_by,
        "remark": remark,
    }
    with mysql.connection_scope() as connection, connection.cursor() as cursor:
        cursor.execute(
            """
            INSERT INTO ip_system_source_snapshot (
              snapshot_id, system_id, source_config_id, snapshot_type,
              raw_snapshot_path, raw_content_hash, raw_metadata_json,
              captured_at, captured_by, remark
            )
            VALUES (
              %(snapshot_id)s, %(system_id)s, %(source_config_id)s, %(snapshot_type)s,
              %(raw_snapshot_path)s, %(raw_content_hash)s, %(raw_metadata_json)s,
              %(captured_at)s, %(captured_by)s, %(remark)s
            )
            """,
            snapshot,
        )
    snapshot["raw_metadata"] = raw_metadata or {}
    snapshot.pop("raw_metadata_json", None)
    return snapshot


def create_sync_batch(
    *,
    system_id: str,
    source_config_id: str | None,
    source_snapshot_id: str | None,
    batch_type: str,
    raw_snapshot_path: str,
    created_by: str,
) -> IpSystemSyncBatch:
    batch = _create_batch_record(
        system_id=system_id,
        source_config_id=source_config_id,
        source_snapshot_id=source_snapshot_id,
        batch_type=batch_type,
        raw_snapshot_path=raw_snapshot_path,
        created_by=created_by,
        status="prepared",
    )
    return IpSystemSyncBatch.model_validate(batch)


def process_import(
    *,
    system_id: str,
    payload: IpSystemImportRequest,
    actor: str,
) -> IpSystemImportResponse:
    system = _get_system(system_id)
    rows, raw_content, parse_errors = _extract_import_rows(payload)
    snapshot = create_source_snapshot(
        system_id=system_id,
        source_config_id=payload.source_config_id,
        snapshot_type=f"{payload.file_format}_import",
        raw_snapshot_path=payload.raw_snapshot_path,
        raw_content=raw_content,
        raw_metadata={
            "import_mode": payload.import_mode,
            "file_format": payload.file_format,
            "row_count": len(rows),
        },
        captured_by=actor,
        remark=payload.remark,
    )
    batch = _create_batch_record(
        system_id=system_id,
        source_config_id=payload.source_config_id,
        source_snapshot_id=str(snapshot["snapshot_id"]),
        batch_type="file_import",
        raw_snapshot_path=payload.raw_snapshot_path,
        created_by=actor,
        status="processing",
    )

    parsed_count = 0
    new_count = 0
    changed_count = 0
    unchanged_count = 0
    exception_count = 0
    removed_count = 0
    seen_keys: set[tuple[str, str, str, str]] = set()
    scopes: set[tuple[str, str]] = set()

    for row_number, row in rows:
        candidate, error = _standardize_candidate(
            row=row,
            row_number=row_number,
            default_system=system,
            source_config_id=payload.source_config_id,
            source_snapshot_id=str(snapshot["snapshot_id"]),
            fallback_source_reference=payload.source_reference,
        )
        if error:
            parse_errors.append(error)
            continue
        parsed_count += 1
        scopes.add((str(candidate["business_domain"]), str(candidate["relation_type_code"])))
        jurisdiction = resolve_candidate_jurisdiction(
            system_id=system_id,
            source_config_id=payload.source_config_id,
            official_name=str(candidate["source_official_name"]),
            official_code=str(candidate["source_official_code"]),
            jurisdiction_name=str(candidate["jurisdiction_name"]),
        )
        if not jurisdiction:
            record_match_exception(
                batch_id=str(batch["batch_id"]),
                system_id=system_id,
                source_config_id=payload.source_config_id,
                official_name=str(candidate["source_official_name"]),
                official_code=str(candidate["source_official_code"]),
                raw_record=candidate,
            )
            exception_count += 1
            continue
        candidate["jurisdiction_id"] = jurisdiction["jurisdiction_id"]
        key = (
            str(candidate["jurisdiction_id"]),
            system_id,
            str(candidate["relation_type_id"]),
            str(candidate["business_domain"]),
        )
        seen_keys.add(key)
        change = _build_change_for_candidate(str(batch["batch_id"]), candidate)
        if change is None:
            unchanged_count += 1
            continue
        if _pending_review_exists(change):
            unchanged_count += 1
            continue
        record_change_review(**change)
        if change["change_type"] == "new_relation":
            new_count += 1
        else:
            changed_count += 1

    if payload.import_mode == "full_snapshot":
        for business_domain, relation_type_code in scopes:
            for change in _build_expire_changes(
                batch_id=str(batch["batch_id"]),
                system_id=system_id,
                business_domain=business_domain,
                relation_type_code=relation_type_code,
                seen_keys=seen_keys,
                source_snapshot_id=str(snapshot["snapshot_id"]),
            ):
                if not _pending_review_exists(change):
                    record_change_review(**change)
                    removed_count += 1

    _finish_batch(
        str(batch["batch_id"]),
        status="completed",
        total=len(rows),
        new_count=new_count,
        changed_count=changed_count,
        removed_count=removed_count,
        unchanged_count=unchanged_count,
        exception_count=exception_count,
        error_message="; ".join(f"row {item.row_number}: {item.reason}" for item in parse_errors[:10]),
    )
    saved_batch = get_sync_batch(str(batch["batch_id"]))
    return IpSystemImportResponse(
        batch=IpSystemSyncBatch.model_validate(saved_batch),
        snapshot_id=str(snapshot["snapshot_id"]),
        total_rows=len(rows),
        parsed_count=parsed_count,
        new_records_count=new_count,
        changed_records_count=changed_count,
        removed_records_count=removed_count,
        unchanged_records_count=unchanged_count,
        exception_records_count=exception_count,
        parse_errors=parse_errors,
    )


def get_sync_batch(batch_id: str) -> IpSystemSyncBatchDetail:
    with mysql.connection_scope() as connection, connection.cursor() as cursor:
        cursor.execute("SELECT * FROM ip_system_sync_batch WHERE batch_id = %s LIMIT 1", (batch_id,))
        batch = cursor.fetchone()
        if not batch:
            raise KeyError(batch_id)
        cursor.execute(
            "SELECT COUNT(*) AS count_value FROM ip_system_change_review WHERE batch_id = %s",
            (batch_id,),
        )
        change_count = int(cursor.fetchone()["count_value"] or 0)
        cursor.execute(
            "SELECT COUNT(*) AS count_value FROM ip_system_match_exception WHERE batch_id = %s",
            (batch_id,),
        )
        exception_count = int(cursor.fetchone()["count_value"] or 0)
    return IpSystemSyncBatchDetail.model_validate(
        {
            **dict(batch),
            "change_review_count": change_count,
            "match_exception_count": exception_count,
        }
    )


def get_batch_changes(batch_id: str) -> list[IpSystemChangeReview]:
    from app.services import ip_system_service

    return ip_system_service.list_change_reviews(batch_id=batch_id)


def get_batch_exceptions(batch_id: str) -> list[dict[str, object]]:
    with mysql.connection_scope() as connection, connection.cursor() as cursor:
        cursor.execute(
            """
            SELECT e.*, s.system_code
            FROM ip_system_match_exception e
            JOIN ip_system_master s ON s.system_id = e.system_id
            WHERE e.batch_id = %s
            ORDER BY e.created_at DESC
            """,
            (batch_id,),
        )
        rows = []
        for row in cursor.fetchall():
            item = dict(row)
            item["raw_record"] = _loads_json_object(item.pop("raw_record_json", None))
            rows.append(item)
    return rows


def reprocess_batch(batch_id: str) -> IpSystemReprocessBatchResponse:
    with mysql.connection_scope() as connection, connection.cursor() as cursor:
        cursor.execute("SELECT * FROM ip_system_sync_batch WHERE batch_id = %s LIMIT 1", (batch_id,))
        batch = cursor.fetchone()
        if not batch:
            raise KeyError(batch_id)
        cursor.execute(
            """
            SELECT *
            FROM ip_system_match_exception
            WHERE batch_id = %s AND status IN ('pending', 'confirmed')
            ORDER BY created_at
            """,
            (batch_id,),
        )
        exceptions = [dict(row) for row in cursor.fetchall()]

    reprocessed = 0
    new_changes = 0
    remaining = 0
    for exception in exceptions:
        raw = _loads_json_object(exception.get("raw_record_json")) or {}
        candidate = dict(raw)
        jurisdiction_id = exception.get("resolved_jurisdiction_id")
        if not jurisdiction_id:
            jurisdiction = resolve_candidate_jurisdiction(
                system_id=str(exception["system_id"]),
                source_config_id=exception.get("source_config_id"),
                official_name=str(exception.get("official_name") or candidate.get("source_official_name") or ""),
                official_code=str(exception.get("official_code") or candidate.get("source_official_code") or ""),
                jurisdiction_name=str(candidate.get("jurisdiction_name") or ""),
            )
            jurisdiction_id = jurisdiction.get("jurisdiction_id") if jurisdiction else None
        if not jurisdiction_id:
            remaining += 1
            continue
        candidate["jurisdiction_id"] = jurisdiction_id
        change = _build_change_for_candidate(batch_id, candidate)
        if change and not _pending_review_exists(change):
            record_change_review(**change)
            new_changes += 1
        _mark_exception_reprocessed(str(exception["exception_id"]), str(jurisdiction_id))
        reprocessed += 1
    _refresh_batch_counts(batch_id)
    return IpSystemReprocessBatchResponse(
        batch_id=batch_id,
        reprocessed_count=reprocessed,
        new_change_count=new_changes,
        remaining_exception_count=remaining,
    )


def normalize_candidate_record(raw: dict[str, object]) -> dict[str, object]:
    official_name = str(raw.get("official_name") or raw.get("jurisdiction_name") or raw.get("name") or "").strip()
    return {
        "official_name": official_name,
        "official_code": str(raw.get("official_code") or raw.get("code") or "").strip().upper(),
        "relation_type_code": str(raw.get("relation_type_code") or "").strip().upper(),
        "business_domain": str(raw.get("business_domain") or "").strip(),
        "effective_date": raw.get("effective_date") or None,
        "expiry_date": raw.get("expiry_date") or None,
        "special_statement": str(raw.get("special_statement") or ""),
        "source_reference": str(raw.get("source_reference") or ""),
        "raw": raw,
    }


def resolve_candidate_jurisdiction(
    *,
    system_id: str,
    source_config_id: str | None,
    official_name: str,
    official_code: str,
    jurisdiction_name: str = "",
) -> dict[str, object] | None:
    names = [item for item in [official_name.strip(), jurisdiction_name.strip()] if item]
    code = official_code.strip().upper()
    with mysql.connection_scope() as connection, connection.cursor() as cursor:
        cursor.execute(
            """
            SELECT m.*, j.jurisdiction_id, j.display_code, j.name_cn, j.name_en
            FROM ip_system_jurisdiction_alias_mapping m
            JOIN jurisdictions j ON j.jurisdiction_id = m.jurisdiction_id
            WHERE m.system_id = %s
              AND m.is_confirmed = 1
              AND ((m.source_config_id <=> %s) OR m.source_config_id IS NULL)
              AND (
                (m.official_code <> '' AND m.official_code = %s)
                OR (m.official_name <> '' AND m.official_name IN %s)
              )
            ORDER BY CASE WHEN m.source_config_id <=> %s THEN 0 ELSE 1 END, m.confirmed_at DESC
            LIMIT 1
            """,
            (system_id, source_config_id, code, tuple(names or [""]), source_config_id),
        )
        alias = cursor.fetchone()
        if alias:
            return dict(alias)
        if code:
            cursor.execute(
                """
                SELECT *
                FROM jurisdictions
                WHERE UPPER(display_code) = %s
                   OR UPPER(internal_code) = %s
                   OR UPPER(COALESCE(wipo_st3_code, '')) = %s
                   OR UPPER(COALESCE(iso_alpha2, '')) = %s
                   OR UPPER(COALESCE(iso_alpha3, '')) = %s
                LIMIT 1
                """,
                (code, code, code, code, code),
            )
            matched = cursor.fetchone()
            if matched:
                return dict(matched)
        for name in names:
            cursor.execute(
                "SELECT * FROM jurisdictions WHERE name_cn = %s OR name_en = %s LIMIT 1",
                (name, name),
            )
            matched = cursor.fetchone()
            if matched:
                return dict(matched)
        cursor.execute("SELECT * FROM jurisdictions")
        normalized_names = {_normalize_name(name) for name in names}
        for row in cursor.fetchall():
            if _normalize_name(str(row.get("name_cn") or "")) in normalized_names:
                return dict(row)
            if _normalize_name(str(row.get("name_en") or "")) in normalized_names:
                return dict(row)
            if _normalize_name(str(row.get("display_code") or "")) in normalized_names:
                return dict(row)
    return None


def record_match_exception(
    *,
    batch_id: str | None,
    system_id: str,
    source_config_id: str | None,
    official_name: str,
    official_code: str,
    raw_record: dict[str, object],
    suggested_jurisdiction_id: str | None = None,
    confidence_score: float | None = None,
) -> dict[str, object]:
    exception = {
        "exception_id": f"ipmex-{uuid4().hex[:16]}",
        "batch_id": batch_id,
        "system_id": system_id,
        "source_config_id": source_config_id,
        "official_name": official_name,
        "official_code": official_code,
        "raw_record_json": json.dumps(raw_record, ensure_ascii=False, default=str),
        "suggested_jurisdiction_id": suggested_jurisdiction_id,
        "confidence_score": confidence_score,
        "status": "pending",
    }
    with mysql.connection_scope() as connection, connection.cursor() as cursor:
        cursor.execute(
            """
            INSERT INTO ip_system_match_exception (
              exception_id, batch_id, system_id, source_config_id, official_name,
              official_code, raw_record_json, suggested_jurisdiction_id,
              confidence_score, status
            )
            VALUES (
              %(exception_id)s, %(batch_id)s, %(system_id)s, %(source_config_id)s,
              %(official_name)s, %(official_code)s, %(raw_record_json)s,
              %(suggested_jurisdiction_id)s, %(confidence_score)s, %(status)s
            )
            """,
            exception,
        )
    exception["raw_record"] = raw_record
    exception.pop("raw_record_json", None)
    return exception


def record_change_review(
    *,
    batch_id: str,
    system_id: str,
    jurisdiction_id: str | None,
    relation_type_id: str | None,
    business_domain: str,
    change_type: str,
    old_relation_id: str | None,
    old_value: dict[str, object] | None,
    new_value: dict[str, object] | None,
    source_official_name: str,
    source_official_code: str,
    data_quality_flags: list[str] | None = None,
    review_comment: str = "",
) -> IpSystemChangeReview:
    review = {
        "review_id": f"iprev-{uuid4().hex[:16]}",
        "batch_id": batch_id,
        "system_id": system_id,
        "jurisdiction_id": jurisdiction_id,
        "relation_type_id": relation_type_id,
        "business_domain": business_domain,
        "change_type": change_type,
        "old_relation_id": old_relation_id,
        "old_value_json": json.dumps(old_value, ensure_ascii=False, default=str),
        "new_value_json": json.dumps(new_value, ensure_ascii=False, default=str),
        "source_official_name": source_official_name,
        "source_official_code": source_official_code,
        "data_quality_flags_json": json.dumps(data_quality_flags or [], ensure_ascii=False),
        "review_status": "pending_review",
        "review_comment": review_comment,
    }
    with mysql.connection_scope() as connection, connection.cursor() as cursor:
        cursor.execute(
            """
            INSERT INTO ip_system_change_review (
              review_id, batch_id, system_id, jurisdiction_id, relation_type_id,
              business_domain, change_type, old_relation_id, old_value_json,
              new_value_json, source_official_name, source_official_code,
              data_quality_flags_json, review_status, review_comment
            )
            VALUES (
              %(review_id)s, %(batch_id)s, %(system_id)s, %(jurisdiction_id)s,
              %(relation_type_id)s, %(business_domain)s, %(change_type)s,
              %(old_relation_id)s, %(old_value_json)s, %(new_value_json)s,
              %(source_official_name)s, %(source_official_code)s,
              %(data_quality_flags_json)s, %(review_status)s, %(review_comment)s
            )
            """,
            review,
        )
    review["old_value"] = old_value
    review["new_value"] = new_value
    review["data_quality_flags"] = data_quality_flags or []
    review["system_code"] = ""
    review["jurisdiction_code"] = ""
    review["jurisdiction_name_cn"] = ""
    review["relation_type_code"] = ""
    review.pop("old_value_json", None)
    review.pop("new_value_json", None)
    review.pop("data_quality_flags_json", None)
    return IpSystemChangeReview.model_validate(review)


def _extract_import_rows(
    payload: IpSystemImportRequest,
) -> tuple[list[tuple[int, dict[str, object]]], str, list[IpSystemImportParseError]]:
    errors: list[IpSystemImportParseError] = []
    if payload.file_format == "rows":
        rows = [(index + 1, row.model_dump()) for index, row in enumerate(payload.rows)]
        raw_content = json.dumps([row for _, row in rows], ensure_ascii=False, default=str)
        return rows, raw_content, errors
    if payload.file_format == "csv":
        reader = csv.DictReader(io.StringIO(payload.csv_content))
        rows = [(index, {key: value for key, value in row.items() if key}) for index, row in enumerate(reader, start=2)]
        return rows, payload.csv_content, errors
    if payload.file_format == "xlsx":
        try:
            workbook = load_workbook(io.BytesIO(base64.b64decode(payload.xlsx_base64)), read_only=True, data_only=True)
        except Exception as exc:
            errors.append(IpSystemImportParseError(row_number=0, reason=f"invalid_xlsx:{exc}", raw={}))
            return [], payload.xlsx_base64, errors
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        try:
            headers = [str(value or "").strip() for value in next(iterator)]
        except StopIteration:
            return [], payload.xlsx_base64, errors
        rows = []
        for row_number, values in enumerate(iterator, start=2):
            if not any(value not in (None, "") for value in values):
                continue
            rows.append((row_number, {headers[index]: value for index, value in enumerate(values) if index < len(headers)}))
        return rows, payload.xlsx_base64, errors
    return [], "", [IpSystemImportParseError(row_number=0, reason="unsupported_file_format", raw={})]


def _standardize_candidate(
    *,
    row: dict[str, object],
    row_number: int,
    default_system: dict[str, object],
    source_config_id: str | None,
    source_snapshot_id: str,
    fallback_source_reference: str,
) -> tuple[dict[str, object], IpSystemImportParseError | None]:
    raw = {key: _clean(value) for key, value in row.items()}
    system_code = str(raw.get("system_code") or default_system["system_code"]).strip().upper()
    if system_code != str(default_system["system_code"]).upper():
        return {}, _row_error(row_number, f"system_code_mismatch:{system_code}", raw)
    business_domain = str(raw.get("business_domain") or "").strip()
    relation_type_code = str(raw.get("relation_type_code") or "").strip().upper()
    if not business_domain:
        return {}, _row_error(row_number, "business_domain_required", raw)
    if not relation_type_code:
        return {}, _row_error(row_number, "relation_type_code_required", raw)
    with mysql.connection_scope() as connection, connection.cursor() as cursor:
        cursor.execute(
            """
            SELECT 1
            FROM ip_system_business_domain
            WHERE system_id = %s AND business_domain = %s AND is_enabled = 1
            LIMIT 1
            """,
            (default_system["system_id"], business_domain),
        )
        if cursor.fetchone() is None:
            return {}, _row_error(row_number, "business_domain_disabled", raw)
        cursor.execute(
            """
            SELECT *
            FROM ip_system_relation_type
            WHERE relation_type_code = %s AND is_active = 1
            LIMIT 1
            """,
            (relation_type_code,),
        )
        relation_type = cursor.fetchone()
        if relation_type is None:
            return {}, _row_error(row_number, "relation_type_not_found", raw)
        source_reference = str(raw.get("source_reference") or fallback_source_reference or "")
        if not source_reference and source_config_id:
            cursor.execute("SELECT source_url FROM ip_system_source_config WHERE source_config_id = %s", (source_config_id,))
            source = cursor.fetchone()
            source_reference = str(source.get("source_url") or "") if source else ""
    effective_date, date_error = _parse_date(raw.get("effective_date"))
    if date_error:
        return {}, _row_error(row_number, f"effective_date_{date_error}", raw)
    expiry_date, date_error = _parse_date(raw.get("expiry_date"))
    if date_error:
        return {}, _row_error(row_number, f"expiry_date_{date_error}", raw)
    if effective_date and expiry_date and expiry_date < effective_date:
        return {}, _row_error(row_number, "expiry_date_before_effective_date", raw)
    jurisdiction_name = str(raw.get("jurisdiction_name") or "").strip()
    official_name = str(raw.get("official_name") or jurisdiction_name).strip()
    official_code = str(raw.get("official_code") or "").strip().upper()
    if not official_name and not official_code and not jurisdiction_name:
        return {}, _row_error(row_number, "jurisdiction_identifier_required", raw)
    flags = _parse_flags(raw.get("data_quality_flags"))
    if effective_date is None and "date_missing" not in flags:
        flags.append("date_missing")
    return {
        "row_number": row_number,
        "system_id": default_system["system_id"],
        "system_code": default_system["system_code"],
        "business_domain": business_domain,
        "relation_type_id": relation_type["relation_type_id"],
        "relation_type_code": relation_type_code,
        "jurisdiction_name": jurisdiction_name,
        "source_official_name": official_name,
        "source_official_code": official_code,
        "effective_date": effective_date,
        "expiry_date": expiry_date,
        "is_active": _parse_bool(raw.get("is_active"), default=True),
        "quote_hint_enabled": _parse_bool(raw.get("quote_hint_enabled"), default=False),
        "path_rule_dependency": _parse_bool(raw.get("path_rule_dependency"), default=False),
        "quote_hint_text": str(raw.get("quote_hint_text") or ""),
        "special_statement": str(raw.get("special_statement") or ""),
        "source_reference": source_reference,
        "source_snapshot_id": source_snapshot_id,
        "data_quality_flags": flags,
        "admin_remark": str(raw.get("admin_remark") or ""),
        "verification_status": "pending_review",
    }, None


def _build_change_for_candidate(batch_id: str, candidate: dict[str, object]) -> dict[str, object] | None:
    existing = _find_published_relation(candidate)
    new_value = _candidate_relation_value(candidate)
    if existing is None:
        return _change_record(batch_id, candidate, "new_relation", None, None, new_value)
    old_value = _relation_value(existing)
    changed_fields = [
        field
        for field in [
            "effective_date",
            "expiry_date",
            "is_active",
            "quote_hint_enabled",
            "path_rule_dependency",
            "quote_hint_text",
            "special_statement",
            "source_reference",
            "source_official_name",
            "source_official_code",
            "data_quality_flags",
            "admin_remark",
        ]
        if _comparable(old_value.get(field)) != _comparable(new_value.get(field))
    ]
    if not changed_fields:
        return None
    if {"effective_date", "expiry_date"} & set(changed_fields):
        change_type = "date_changed"
    elif {"source_reference", "source_official_name", "source_official_code"} & set(changed_fields):
        change_type = "source_changed"
    elif {"special_statement", "admin_remark"} & set(changed_fields):
        change_type = "remark_changed"
    else:
        change_type = "update_relation"
    new_value["relation_id"] = existing["relation_id"]
    return _change_record(batch_id, candidate, change_type, existing["relation_id"], old_value, new_value)


def _build_expire_changes(
    *,
    batch_id: str,
    system_id: str,
    business_domain: str,
    relation_type_code: str,
    seen_keys: set[tuple[str, str, str, str]],
    source_snapshot_id: str,
) -> list[dict[str, object]]:
    changes: list[dict[str, object]] = []
    with mysql.connection_scope() as connection, connection.cursor() as cursor:
        cursor.execute(
            """
            SELECT r.*, rt.relation_type_code
            FROM jurisdiction_ip_system_relation r
            JOIN ip_system_relation_type rt ON rt.relation_type_id = r.relation_type_id
            WHERE r.system_id = %s
              AND r.business_domain = %s
              AND rt.relation_type_code = %s
              AND r.publish_status = 'published'
              AND r.is_active = 1
            """,
            (system_id, business_domain, relation_type_code),
        )
        rows = [dict(row) for row in cursor.fetchall()]
    for row in rows:
        key = (row["jurisdiction_id"], row["system_id"], row["relation_type_id"], row["business_domain"])
        if key in seen_keys:
            continue
        old_value = _relation_value(row)
        new_value = {
            **old_value,
            "is_active": False,
            "expiry_date": date.today(),
            "source_snapshot_id": source_snapshot_id,
        }
        candidate = {
            "system_id": row["system_id"],
            "jurisdiction_id": row["jurisdiction_id"],
            "relation_type_id": row["relation_type_id"],
            "business_domain": row["business_domain"],
            "source_official_name": row.get("source_official_name") or "",
            "source_official_code": row.get("source_official_code") or "",
            "data_quality_flags": _loads_json_list(row.get("data_quality_flags_json")),
        }
        changes.append(_change_record(batch_id, candidate, "expire_relation", row["relation_id"], old_value, new_value))
    return changes


def _find_published_relation(candidate: dict[str, object]) -> dict[str, object] | None:
    with mysql.connection_scope() as connection, connection.cursor() as cursor:
        cursor.execute(
            """
            SELECT *
            FROM jurisdiction_ip_system_relation
            WHERE jurisdiction_id = %s
              AND system_id = %s
              AND relation_type_id = %s
              AND business_domain = %s
              AND publish_status = 'published'
            ORDER BY is_active DESC, COALESCE(expiry_date, DATE('9999-12-31')) DESC, created_at DESC
            LIMIT 1
            """,
            (
                candidate["jurisdiction_id"],
                candidate["system_id"],
                candidate["relation_type_id"],
                candidate["business_domain"],
            ),
        )
        row = cursor.fetchone()
    return dict(row) if row else None


def _pending_review_exists(change: dict[str, object]) -> bool:
    new_json = json.dumps(
        _dedupe_value(change.get("new_value")),
        ensure_ascii=False,
        default=str,
        sort_keys=True,
    )
    with mysql.connection_scope() as connection, connection.cursor() as cursor:
        cursor.execute(
            """
            SELECT new_value_json
            FROM ip_system_change_review
            WHERE system_id = %s
              AND (jurisdiction_id <=> %s)
              AND (relation_type_id <=> %s)
              AND business_domain = %s
              AND change_type = %s
              AND review_status = 'pending_review'
            """,
            (
                change["system_id"],
                change["jurisdiction_id"],
                change["relation_type_id"],
                change["business_domain"],
                change["change_type"],
            ),
        )
        for row in cursor.fetchall():
            existing_json = json.dumps(
                _dedupe_value(_loads_json_object(row["new_value_json"])),
                ensure_ascii=False,
                default=str,
                sort_keys=True,
            )
            if existing_json == new_json:
                return True
    return False


def _dedupe_value(value: object) -> object:
    if not isinstance(value, dict):
        return value
    normalized = dict(value)
    normalized.pop("source_snapshot_id", None)
    return normalized


def _change_record(
    batch_id: str,
    candidate: dict[str, object],
    change_type: str,
    old_relation_id: str | None,
    old_value: dict[str, object] | None,
    new_value: dict[str, object] | None,
) -> dict[str, object]:
    return {
        "batch_id": batch_id,
        "system_id": str(candidate["system_id"]),
        "jurisdiction_id": candidate.get("jurisdiction_id"),
        "relation_type_id": candidate.get("relation_type_id"),
        "business_domain": str(candidate.get("business_domain") or ""),
        "change_type": change_type,
        "old_relation_id": old_relation_id,
        "old_value": old_value,
        "new_value": new_value,
        "source_official_name": str(candidate.get("source_official_name") or ""),
        "source_official_code": str(candidate.get("source_official_code") or ""),
        "data_quality_flags": list(candidate.get("data_quality_flags") or []),
        "review_comment": "",
    }


def _candidate_relation_value(candidate: dict[str, object]) -> dict[str, object]:
    return {
        "jurisdiction_id": candidate.get("jurisdiction_id"),
        "system_id": candidate.get("system_id"),
        "relation_type_id": candidate.get("relation_type_id"),
        "business_domain": candidate.get("business_domain"),
        "is_active": candidate.get("is_active", True),
        "effective_date": candidate.get("effective_date"),
        "expiry_date": candidate.get("expiry_date"),
        "source_reference": candidate.get("source_reference", ""),
        "source_snapshot_id": candidate.get("source_snapshot_id"),
        "source_official_name": candidate.get("source_official_name", ""),
        "source_official_code": candidate.get("source_official_code", ""),
        "verification_status": candidate.get("verification_status", "pending_review"),
        "quote_hint_enabled": candidate.get("quote_hint_enabled", False),
        "path_rule_dependency": candidate.get("path_rule_dependency", False),
        "quote_hint_text": candidate.get("quote_hint_text", ""),
        "special_statement": candidate.get("special_statement", ""),
        "data_quality_flags": list(candidate.get("data_quality_flags") or []),
        "admin_remark": candidate.get("admin_remark", ""),
    }


def _relation_value(row: dict[str, object]) -> dict[str, object]:
    return {
        "relation_id": row.get("relation_id"),
        "jurisdiction_id": row.get("jurisdiction_id"),
        "system_id": row.get("system_id"),
        "relation_type_id": row.get("relation_type_id"),
        "business_domain": row.get("business_domain"),
        "is_active": bool(row.get("is_active")),
        "effective_date": row.get("effective_date"),
        "expiry_date": row.get("expiry_date"),
        "source_reference": row.get("source_reference") or "",
        "source_snapshot_id": row.get("source_snapshot_id"),
        "source_official_name": row.get("source_official_name") or "",
        "source_official_code": row.get("source_official_code") or "",
        "verification_status": row.get("verification_status") or "",
        "quote_hint_enabled": bool(row.get("quote_hint_enabled")),
        "path_rule_dependency": bool(row.get("path_rule_dependency")),
        "quote_hint_text": row.get("quote_hint_text") or "",
        "special_statement": row.get("special_statement") or "",
        "data_quality_flags": _loads_json_list(row.get("data_quality_flags_json")),
        "admin_remark": row.get("admin_remark") or "",
    }


def _create_batch_record(
    *,
    system_id: str,
    source_config_id: str | None,
    source_snapshot_id: str | None,
    batch_type: str,
    raw_snapshot_path: str,
    created_by: str,
    status: str,
) -> dict[str, object]:
    batch = {
        "batch_id": f"ipsync-{uuid4().hex[:16]}",
        "system_id": system_id,
        "source_config_id": source_config_id,
        "source_snapshot_id": source_snapshot_id,
        "batch_type": batch_type,
        "started_at": datetime.now(),
        "status": status,
        "raw_snapshot_path": raw_snapshot_path,
        "created_by": created_by,
    }
    with mysql.connection_scope() as connection, connection.cursor() as cursor:
        cursor.execute(
            """
            INSERT INTO ip_system_sync_batch (
              batch_id, system_id, source_config_id, source_snapshot_id, batch_type,
              started_at, status, raw_snapshot_path, created_by
            )
            VALUES (
              %(batch_id)s, %(system_id)s, %(source_config_id)s, %(source_snapshot_id)s,
              %(batch_type)s, %(started_at)s, %(status)s, %(raw_snapshot_path)s,
              %(created_by)s
            )
            """,
            batch,
        )
        cursor.execute("SELECT * FROM ip_system_sync_batch WHERE batch_id = %s", (batch["batch_id"],))
        return dict(cursor.fetchone())


def _finish_batch(
    batch_id: str,
    *,
    status: str,
    total: int,
    new_count: int,
    changed_count: int,
    removed_count: int,
    unchanged_count: int,
    exception_count: int,
    error_message: str,
) -> None:
    with mysql.connection_scope() as connection, connection.cursor() as cursor:
        cursor.execute(
            """
            UPDATE ip_system_sync_batch
            SET finished_at = %s,
                status = %s,
                total_records_found = %s,
                new_records_count = %s,
                changed_records_count = %s,
                removed_records_count = %s,
                unchanged_records_count = %s,
                exception_records_count = %s,
                error_message = %s
            WHERE batch_id = %s
            """,
            (
                datetime.now(),
                status,
                total,
                new_count,
                changed_count,
                removed_count,
                unchanged_count,
                exception_count,
                error_message,
                batch_id,
            ),
        )


def _refresh_batch_counts(batch_id: str) -> None:
    with mysql.connection_scope() as connection, connection.cursor() as cursor:
        cursor.execute(
            "SELECT COUNT(*) AS count_value FROM ip_system_change_review WHERE batch_id = %s",
            (batch_id,),
        )
        change_count = int(cursor.fetchone()["count_value"] or 0)
        cursor.execute(
            "SELECT COUNT(*) AS count_value FROM ip_system_match_exception WHERE batch_id = %s AND status = 'pending'",
            (batch_id,),
        )
        exception_count = int(cursor.fetchone()["count_value"] or 0)
        cursor.execute(
            """
            UPDATE ip_system_sync_batch
            SET changed_records_count = %s,
                exception_records_count = %s,
                status = CASE WHEN %s = 0 THEN 'completed' ELSE status END
            WHERE batch_id = %s
            """,
            (change_count, exception_count, exception_count, batch_id),
        )


def _mark_exception_reprocessed(exception_id: str, jurisdiction_id: str) -> None:
    with mysql.connection_scope() as connection, connection.cursor() as cursor:
        cursor.execute(
            """
            UPDATE ip_system_match_exception
            SET status = 'reprocessed',
                resolved_jurisdiction_id = COALESCE(resolved_jurisdiction_id, %s),
                resolved_at = COALESCE(resolved_at, %s)
            WHERE exception_id = %s
            """,
            (jurisdiction_id, datetime.now(), exception_id),
        )


def _get_system(system_id: str) -> dict[str, object]:
    with mysql.connection_scope() as connection, connection.cursor() as cursor:
        cursor.execute("SELECT * FROM ip_system_master WHERE system_id = %s LIMIT 1", (system_id,))
        system = cursor.fetchone()
    if not system:
        raise KeyError(system_id)
    return dict(system)


def _parse_date(value: object) -> tuple[date | None, str | None]:
    if value in (None, ""):
        return None, None
    if isinstance(value, datetime):
        return value.date(), None
    if isinstance(value, date):
        return value, None
    text = str(value).strip()
    if not text:
        return None, None
    for candidate in [text, text.replace("/", "-"), text.replace(".", "-")]:
        try:
            return date.fromisoformat(candidate[:10]), None
        except ValueError:
            continue
    return None, "invalid"


def _parse_bool(value: object, default: bool) -> bool:
    if value in (None, ""):
        return default
    if isinstance(value, bool):
        return value
    if isinstance(value, int):
        return bool(value)
    text = str(value).strip().lower()
    if text in {"1", "true", "yes", "y", "是", "启用", "有效"}:
        return True
    if text in {"0", "false", "no", "n", "否", "停用", "无效"}:
        return False
    return default


def _parse_flags(value: object) -> list[str]:
    if value in (None, ""):
        return []
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    text = str(value).strip()
    if not text:
        return []
    try:
        loaded = json.loads(text)
    except json.JSONDecodeError:
        loaded = None
    if isinstance(loaded, list):
        return [str(item).strip() for item in loaded if str(item).strip()]
    return [item.strip() for item in re.split(r"[,，;；]", text) if item.strip()]


def _loads_json_list(value: object) -> list[str]:
    if value in (None, ""):
        return []
    if isinstance(value, list):
        return [str(item) for item in value if str(item)]
    try:
        loaded = json.loads(str(value))
    except json.JSONDecodeError:
        return []
    if not isinstance(loaded, list):
        return []
    return [str(item) for item in loaded if str(item)]


def _loads_json_object(value: object) -> dict[str, object] | None:
    if value in (None, ""):
        return None
    if isinstance(value, dict):
        return value
    try:
        loaded = json.loads(str(value))
    except json.JSONDecodeError:
        return None
    return loaded if isinstance(loaded, dict) else None


def _row_error(row_number: int, reason: str, raw: dict[str, object]) -> IpSystemImportParseError:
    return IpSystemImportParseError(row_number=row_number, reason=reason, raw=raw)


def _clean(value: object) -> object:
    if isinstance(value, str):
        return value.strip()
    return value


def _normalize_name(value: str) -> str:
    return re.sub(r"[\s\W_]+", "", value.lower(), flags=re.UNICODE)


def _comparable(value: object) -> object:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, list):
        return sorted(str(item) for item in value)
    return value
